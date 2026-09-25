from docx import Document 
from datetime import date 
 
 
def generate_weekly_report(metrics: dict, narrative: str, output_path: str): 
   doc = Document() 
   doc.add_heading("Weekly Sales Report", level=1) 
   doc.add_paragraph(f"Report date: {date.today().strftime('%B %d, %Y')}") 
 
   doc.add_heading("Summary", level=2) 
   doc.add_paragraph(narrative) 
 
   doc.add_heading("Key Metrics", level=2) 
   table = doc.add_table(rows=1, cols=2) 
   table.style = "Light Grid Accent 1" 
 
   hdr = table.rows[0].cells 
   hdr[0].text, hdr[1].text = "Metric", "Value" 
 
   for metric, value in metrics.items(): 
       row = table.add_row().cells 
       row[0].text, row[1].text = metric, str(value) 
 
   doc.save(output_path) 
